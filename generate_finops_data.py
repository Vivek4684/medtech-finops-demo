#!/usr/bin/env python3
"""
FinOps Data Platform cost-model generator (dummy data).

Produces:
  - data/finops_data.xlsx   -> the "source of truth" workbook for the Chief
  - data/finops_data.json   -> embedded into the standalone HTML dashboard

Model
-----
Total annual spend ~ $20,000,000 across the data-platform tech stack.

Each TECHNOLOGY has spend split into:
  * Shared / Enterprise  (enterprise licenses + platform base fees -> a shared pool)
  * Sector-attributed    (direct consumption, split across the 4 sectors)

The Shared pool is ALSO allocated to sectors via an allocation key (headcount/usage
weighted) so we can show a "fully-loaded" cost per sector (a FinOps hallmark).

Sector-attributed spend drills down to PROJECTS.
Every level carries an "Other / Everything Else" catch-all so totals reconcile.
"""

import json
import os
import random
from collections import defaultdict

random.seed(42)

OUT_DIR = os.path.join(os.path.dirname(__file__), "data")
os.makedirs(OUT_DIR, exist_ok=True)

TARGET_TOTAL = 20_000_000

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

SECTORS = ["Supply Chain", "Commercial", "R&D", "Enterprise"]

# ---------------------------------------------------------------------------
# 1. Technology stack: annual $ and split between shared pool vs attributable.
#    "shared_frac" = portion that is enterprise license / platform base fee.
#    Numbers chosen to be realistic (AWS + Databricks dominate) and to sum to 20M.
# ---------------------------------------------------------------------------
TECHS = [
    # name,            annual$,    shared_frac, category
    ("AWS Cloud",       5_200_000, 0.25, "Cloud Infrastructure"),
    ("Databricks",      4_300_000, 0.30, "Data & Analytics Platform"),
    ("Snowflake",       1_900_000, 0.35, "Data & Analytics Platform"),
    ("Informatica",     1_450_000, 0.60, "Data Integration"),
    ("Power BI",        1_150_000, 0.70, "BI & Visualization"),
    ("Tableau",           950_000, 0.75, "BI & Visualization"),
    ("Confluent Kafka",   900_000, 0.45, "Streaming / Eventing"),
    ("Fivetran",          750_000, 0.50, "Data Integration"),
    ("Denodo",            700_000, 0.65, "Data Virtualization"),
    ("Collibra",          600_000, 0.80, "Data Governance"),
    ("dbt Cloud",         350_000, 0.55, "Transformation"),
    ("Other / Everything Else", 1_750_000, 0.40, "Misc SaaS & Services"),
]

assert sum(t[1] for t in TECHS) == TARGET_TOTAL, sum(t[1] for t in TECHS)

# ---------------------------------------------------------------------------
# 2. Projects and their sector.  Each sector also gets an "Other Projects" line.
# ---------------------------------------------------------------------------
PROJECTS = [
    # name,              sector,         weight (relative size within its sector)
    ("CDL",              "Supply Chain", 0.70),
    ("Radar",            "R&D",          0.72),
    ("STAR",             "Commercial",   0.34),
    ("Datavantage",      "Commercial",   0.30),
    ("Artemis / Hubble", "Commercial",   0.24),
]

# Sector-attributed spend split across the 4 sectors (consumption weighting).
# Each tech can lean differently; we give a base profile then jitter per tech.
SECTOR_BASE_WEIGHT = {
    "Supply Chain": 0.26,
    "Commercial":   0.34,
    "R&D":          0.28,
    "Enterprise":   0.12,  # enterprise-run shared workloads not tied to a sector
}

# Per-sector share that lands on named projects vs "Other Projects".
NAMED_PROJECT_COVERAGE = {
    "Supply Chain": 0.82,  # CDL dominates
    "R&D":          0.80,  # Radar dominates
    "Commercial":   0.85,  # STAR + Datavantage + Artemis/Hubble
    "Enterprise":   0.0,   # enterprise = shared, no named demo projects
}


def seasonal_monthly_split(annual, growth=0.0):
    """Split an annual figure into 12 months with mild seasonality + slight growth."""
    base = [1.0 + 0.04 * random.uniform(-1, 1) for _ in MONTHS]
    # gentle upward ramp over the year
    base = [b * (1 + growth * (i / 11.0)) for i, b in enumerate(base)]
    s = sum(base)
    return [round(annual * b / s, 2) for b in base]


def jitter_weights(weights):
    out = {k: w * (1 + 0.15 * random.uniform(-1, 1)) for k, w in weights.items()}
    s = sum(out.values())
    return {k: v / s for k, v in out.items()}


# ---------------------------------------------------------------------------
# Build the normalized fact table: one row per (tech, bucket, sector, project, month)
# ---------------------------------------------------------------------------
rows = []  # flat fact rows

for name, annual, shared_frac, category in TECHS:
    shared_annual = round(annual * shared_frac)
    attributable_annual = annual - shared_annual

    # ---- Shared / Enterprise pool (licenses + platform) -------------------
    monthly = seasonal_monthly_split(shared_annual)
    for i, m in enumerate(MONTHS):
        rows.append({
            "Technology": name,
            "Category": category,
            "CostType": "Shared / Enterprise",
            "Sector": "Enterprise (Shared Pool)",
            "Project": "Platform & Licenses",
            "Month": m,
            "Amount": monthly[i],
        })

    # ---- Sector-attributed consumption -----------------------------------
    sw = jitter_weights(SECTOR_BASE_WEIGHT)
    for sector in SECTORS:
        sector_annual = round(attributable_annual * sw[sector])
        if sector_annual <= 0:
            continue

        # projects within this sector
        coverage = NAMED_PROJECT_COVERAGE[sector]
        sector_projects = [(p, w) for (p, s, w) in PROJECTS if s == sector]

        alloc = []  # (project_name, amount)
        if sector_projects and coverage > 0:
            named_pool = round(sector_annual * coverage)
            wsum = sum(w for _, w in sector_projects)
            assigned = 0
            for p, w in sector_projects:
                amt = round(named_pool * w / wsum)
                alloc.append((p, amt))
                assigned += amt
            # remainder -> Other Projects
            other = sector_annual - assigned
            if other > 0:
                alloc.append((f"Other {sector} Projects", other))
        else:
            alloc.append(("Enterprise Workloads" if sector == "Enterprise"
                          else f"Other {sector} Projects", sector_annual))

        for project, amt in alloc:
            monthly = seasonal_monthly_split(amt, growth=random.uniform(0.0, 0.25))
            for i, m in enumerate(MONTHS):
                rows.append({
                    "Technology": name,
                    "Category": category,
                    "CostType": "Sector-Attributed",
                    "Sector": sector,
                    "Project": project,
                    "Month": m,
                    "Amount": monthly[i],
                })

# ---------------------------------------------------------------------------
# Allocation of the shared pool to sectors (fully-loaded view).
# Allocation key: weighted by each sector's attributable consumption + a base.
# ---------------------------------------------------------------------------
attributable_by_sector = defaultdict(float)
for r in rows:
    if r["CostType"] == "Sector-Attributed":
        attributable_by_sector[r["Sector"]] += r["Amount"]

total_shared = sum(r["Amount"] for r in rows if r["CostType"] == "Shared / Enterprise")
# allocation weights: 70% by consumption, 30% split evenly (platform availability)
alloc_sectors = SECTORS
consume_total = sum(attributable_by_sector[s] for s in alloc_sectors)
alloc_weights = {}
for s in alloc_sectors:
    alloc_weights[s] = 0.70 * (attributable_by_sector[s] / consume_total) + 0.30 * (1 / len(alloc_sectors))

shared_alloc = {s: round(total_shared * alloc_weights[s]) for s in alloc_sectors}

# ---------------------------------------------------------------------------
# Aggregations for the dashboard JSON
# ---------------------------------------------------------------------------
def rnd(x):
    return round(float(x), 2)

total_spend = rnd(sum(r["Amount"] for r in rows))

# by technology
tech_summary = {}
for name, annual, shared_frac, category in TECHS:
    tr = [r for r in rows if r["Technology"] == name]
    shared = rnd(sum(r["Amount"] for r in tr if r["CostType"] == "Shared / Enterprise"))
    attributable = rnd(sum(r["Amount"] for r in tr if r["CostType"] == "Sector-Attributed"))
    by_sector = defaultdict(float)
    for r in tr:
        if r["CostType"] == "Sector-Attributed":
            by_sector[r["Sector"]] += r["Amount"]
    by_project = defaultdict(float)
    for r in tr:
        if r["CostType"] == "Sector-Attributed":
            by_project[r["Project"]] += r["Amount"]
    monthly = defaultdict(float)
    for r in tr:
        monthly[r["Month"]] += r["Amount"]
    tech_summary[name] = {
        "category": category,
        "total": rnd(shared + attributable),
        "shared": shared,
        "attributable": attributable,
        "by_sector": {k: rnd(v) for k, v in by_sector.items()},
        "by_project": {k: rnd(v) for k, v in sorted(by_project.items(), key=lambda x: -x[1])},
        "monthly": [rnd(monthly[m]) for m in MONTHS],
    }

# by sector (direct + allocated shared = fully loaded)
sector_summary = {}
for s in SECTORS:
    direct = rnd(attributable_by_sector[s])
    allocated = rnd(shared_alloc[s])
    # tech breakdown within the sector (direct)
    by_tech = defaultdict(float)
    by_project = defaultdict(float)
    for r in rows:
        if r["CostType"] == "Sector-Attributed" and r["Sector"] == s:
            by_tech[r["Technology"]] += r["Amount"]
            by_project[r["Project"]] += r["Amount"]
    sector_summary[s] = {
        "direct": direct,
        "allocated_shared": allocated,
        "fully_loaded": rnd(direct + allocated),
        "by_tech": {k: rnd(v) for k, v in sorted(by_tech.items(), key=lambda x: -x[1])},
        "by_project": {k: rnd(v) for k, v in sorted(by_project.items(), key=lambda x: -x[1])},
    }

# project -> (sector, tech breakdown)
project_summary = {}
proj_sector = {p: s for (p, s, w) in PROJECTS}
for r in rows:
    if r["CostType"] != "Sector-Attributed":
        continue
    p = r["Project"]
    if p not in project_summary:
        project_summary[p] = {
            "sector": r["Sector"],
            "total": 0.0,
            "by_tech": defaultdict(float),
            "monthly": defaultdict(float),
        }
    project_summary[p]["total"] += r["Amount"]
    project_summary[p]["by_tech"][r["Technology"]] += r["Amount"]
    project_summary[p]["monthly"][r["Month"]] += r["Amount"]

for p, d in project_summary.items():
    d["total"] = rnd(d["total"])
    d["by_tech"] = {k: rnd(v) for k, v in sorted(d["by_tech"].items(), key=lambda x: -x[1])}
    d["monthly"] = [rnd(d["monthly"][m]) for m in MONTHS]

# monthly totals (enterprise)
monthly_total = defaultdict(float)
for r in rows:
    monthly_total[r["Month"]] += r["Amount"]

shared_total = rnd(sum(r["Amount"] for r in rows if r["CostType"] == "Shared / Enterprise"))
attributable_total = rnd(sum(r["Amount"] for r in rows if r["CostType"] == "Sector-Attributed"))

dashboard = {
    "meta": {
        "company": "MedTech Data Platform — FinOps Demo",
        "currency": "USD",
        "total_spend": total_spend,
        "shared_total": shared_total,
        "attributable_total": attributable_total,
        "months": MONTHS,
        "sectors": SECTORS,
    },
    "monthly_total": [rnd(monthly_total[m]) for m in MONTHS],
    "technologies": tech_summary,
    "sectors": sector_summary,
    "projects": project_summary,
    "shared_allocation": {k: rnd(v) for k, v in shared_alloc.items()},
}

with open(os.path.join(OUT_DIR, "finops_data.json"), "w") as f:
    json.dump(dashboard, f, indent=2)

print(f"Total spend: ${total_spend:,.0f}")
print(f"  Shared/Enterprise pool: ${shared_total:,.0f}")
print(f"  Sector-attributed:      ${attributable_total:,.0f}")
print(f"Fact rows: {len(rows)}")

# ===========================================================================
# Excel workbook
# ===========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, PieChart, LineChart, Reference

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="2E5496")
ACCENT_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="BDD7EE")
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0'
CENTER = Alignment(horizontal="center", vertical="center")

wb = Workbook()

def style_header(ws, row, ncols, fill=HEADER_FILL):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = WHITE
        cell.alignment = CENTER
        cell.border = BORDER

def autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

# ---- Sheet 1: Executive Summary ------------------------------------------
ws = wb.active
ws.title = "Executive Summary"
ws["A1"] = "MedTech Data Platform — FinOps Executive Summary"
ws["A1"].font = Font(size=16, bold=True, color="1F3864")
ws["A2"] = "Dummy data for demonstration. Annual view, USD."
ws["A2"].font = Font(italic=True, color="808080")

ws["A4"] = "Total Annual Spend"
ws["B4"] = total_spend
ws["A5"] = "Shared / Enterprise (licenses + platform)"
ws["B5"] = shared_total
ws["A6"] = "Sector-Attributed (direct consumption)"
ws["B6"] = attributable_total
for r in (4, 5, 6):
    ws.cell(row=r, column=1).font = BOLD
    ws.cell(row=r, column=2).number_format = MONEY
    ws.cell(row=r, column=2).font = BOLD

ws["A8"] = "Fully-Loaded Cost by Sector (direct + allocated shared)"
ws["A8"].font = Font(bold=True, color="1F3864")
hdr = ["Sector", "Direct", "Allocated Shared", "Fully-Loaded", "% of Total"]
ws.append([])  # row 9 placeholder handled below
for j, h in enumerate(hdr, start=1):
    ws.cell(row=9, column=j, value=h)
style_header(ws, 9, len(hdr))
row = 10
for s in SECTORS:
    d = sector_summary[s]
    ws.cell(row=row, column=1, value=s)
    ws.cell(row=row, column=2, value=d["direct"]).number_format = MONEY
    ws.cell(row=row, column=3, value=d["allocated_shared"]).number_format = MONEY
    ws.cell(row=row, column=4, value=d["fully_loaded"]).number_format = MONEY
    ws.cell(row=row, column=5, value=d["fully_loaded"] / total_spend).number_format = "0.0%"
    for c in range(1, 6):
        ws.cell(row=row, column=c).border = BORDER
    row += 1
ws.cell(row=row, column=1, value="TOTAL").font = BOLD
ws.cell(row=row, column=4, value=f"=SUM(D10:D{row-1})").number_format = MONEY
ws.cell(row=row, column=4).font = BOLD
for c in range(1, 6):
    ws.cell(row=row, column=c).fill = TOTAL_FILL
autofit(ws, [40, 16, 18, 16, 12])

# Pie chart of fully-loaded by sector
pie = PieChart()
pie.title = "Fully-Loaded Cost by Sector"
labels = Reference(ws, min_col=1, min_row=10, max_row=10 + len(SECTORS) - 1)
data = Reference(ws, min_col=4, min_row=9, max_row=10 + len(SECTORS) - 1)
pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)
pie.height, pie.width = 8, 12
ws.add_chart(pie, "G9")

# ---- Sheet 2: Spend by Technology ----------------------------------------
ws2 = wb.create_sheet("Spend by Technology")
ws2["A1"] = "Annual Spend by Technology"
ws2["A1"].font = Font(size=14, bold=True, color="1F3864")
hdr = ["Technology", "Category", "Shared / Enterprise", "Sector-Attributed", "Total", "% of Total"]
for j, h in enumerate(hdr, start=1):
    ws2.cell(row=3, column=j, value=h)
style_header(ws2, 3, len(hdr))
row = 4
for name, annual, shared_frac, category in TECHS:
    t = tech_summary[name]
    ws2.cell(row=row, column=1, value=name)
    ws2.cell(row=row, column=2, value=category)
    ws2.cell(row=row, column=3, value=t["shared"]).number_format = MONEY
    ws2.cell(row=row, column=4, value=t["attributable"]).number_format = MONEY
    ws2.cell(row=row, column=5, value=t["total"]).number_format = MONEY
    ws2.cell(row=row, column=6, value=t["total"] / total_spend).number_format = "0.0%"
    for c in range(1, 7):
        ws2.cell(row=row, column=c).border = BORDER
    row += 1
ws2.cell(row=row, column=1, value="TOTAL").font = BOLD
ws2.cell(row=row, column=5, value=f"=SUM(E4:E{row-1})").number_format = MONEY
ws2.cell(row=row, column=5).font = BOLD
for c in range(1, 7):
    ws2.cell(row=row, column=c).fill = TOTAL_FILL
autofit(ws2, [26, 26, 20, 18, 16, 12])

bar = BarChart()
bar.type = "bar"
bar.title = "Total Spend by Technology"
data = Reference(ws2, min_col=5, min_row=3, max_row=3 + len(TECHS))
cats = Reference(ws2, min_col=1, min_row=4, max_row=3 + len(TECHS))
bar.add_data(data, titles_from_data=True)
bar.set_categories(cats)
bar.height, bar.width = 12, 16
ws2.add_chart(bar, "H3")

# ---- Sheet 3: Technology x Sector matrix ---------------------------------
ws3 = wb.create_sheet("Tech x Sector Matrix")
ws3["A1"] = "Sector-Attributed Spend: Technology x Sector"
ws3["A1"].font = Font(size=14, bold=True, color="1F3864")
hdr = ["Technology"] + SECTORS + ["Total"]
for j, h in enumerate(hdr, start=1):
    ws3.cell(row=3, column=j, value=h)
style_header(ws3, 3, len(hdr))
row = 4
for name, annual, shared_frac, category in TECHS:
    t = tech_summary[name]
    ws3.cell(row=row, column=1, value=name)
    for j, s in enumerate(SECTORS, start=2):
        ws3.cell(row=row, column=j, value=t["by_sector"].get(s, 0)).number_format = MONEY
    ws3.cell(row=row, column=2 + len(SECTORS),
             value=f"=SUM(B{row}:{get_column_letter(1+len(SECTORS))}{row})").number_format = MONEY
    for c in range(1, len(hdr) + 1):
        ws3.cell(row=row, column=c).border = BORDER
    row += 1
ws3.cell(row=row, column=1, value="TOTAL").font = BOLD
for j in range(2, len(hdr) + 1):
    col = get_column_letter(j)
    ws3.cell(row=row, column=j, value=f"=SUM({col}4:{col}{row-1})").number_format = MONEY
    ws3.cell(row=row, column=j).font = BOLD
for c in range(1, len(hdr) + 1):
    ws3.cell(row=row, column=c).fill = TOTAL_FILL
autofit(ws3, [26, 16, 16, 14, 16, 16])

# ---- Sheet 4: Spend by Project -------------------------------------------
ws4 = wb.create_sheet("Spend by Project")
ws4["A1"] = "Annual Spend by Project (drill-down)"
ws4["A1"].font = Font(size=14, bold=True, color="1F3864")
hdr = ["Project", "Sector", "Total", "Top Technology", "% of Total"]
for j, h in enumerate(hdr, start=1):
    ws4.cell(row=3, column=j, value=h)
style_header(ws4, 3, len(hdr))
row = 4
# order: named projects first, then "Other ..." lines
ordered = sorted(project_summary.items(),
                 key=lambda kv: (kv[0].startswith("Other"), -kv[1]["total"]))
for p, d in ordered:
    top_tech = next(iter(d["by_tech"]), "-")
    ws4.cell(row=row, column=1, value=p)
    ws4.cell(row=row, column=2, value=d["sector"])
    ws4.cell(row=row, column=3, value=d["total"]).number_format = MONEY
    ws4.cell(row=row, column=4, value=top_tech)
    ws4.cell(row=row, column=5, value=d["total"] / total_spend).number_format = "0.0%"
    for c in range(1, 6):
        ws4.cell(row=row, column=c).border = BORDER
    row += 1
autofit(ws4, [26, 16, 16, 20, 12])

# ---- Sheet 5: Monthly Trend ----------------------------------------------
ws5 = wb.create_sheet("Monthly Trend")
ws5["A1"] = "Monthly Spend Trend (All Technologies)"
ws5["A1"].font = Font(size=14, bold=True, color="1F3864")
ws5.cell(row=3, column=1, value="Month")
ws5.cell(row=3, column=2, value="Total Spend")
style_header(ws5, 3, 2)
for i, m in enumerate(MONTHS):
    ws5.cell(row=4 + i, column=1, value=m)
    ws5.cell(row=4 + i, column=2, value=dashboard["monthly_total"][i]).number_format = MONEY
    ws5.cell(row=4 + i, column=1).border = BORDER
    ws5.cell(row=4 + i, column=2).border = BORDER
autofit(ws5, [14, 18])
line = LineChart()
line.title = "Monthly Spend Trend"
data = Reference(ws5, min_col=2, min_row=3, max_row=3 + len(MONTHS))
cats = Reference(ws5, min_col=1, min_row=4, max_row=3 + len(MONTHS))
line.add_data(data, titles_from_data=True)
line.set_categories(cats)
line.height, line.width = 10, 18
ws5.add_chart(line, "D3")

# ---- Sheet 6: Raw Fact Table ---------------------------------------------
ws6 = wb.create_sheet("Raw Data")
cols = ["Technology", "Category", "CostType", "Sector", "Project", "Month", "Amount"]
ws6.append(cols)
style_header(ws6, 1, len(cols))
for r in rows:
    ws6.append([r[c] for c in cols])
for rr in range(2, len(rows) + 2):
    ws6.cell(row=rr, column=7).number_format = MONEY
autofit(ws6, [26, 26, 20, 22, 26, 8, 14])
ws6.freeze_panes = "A2"

xlsx_path = os.path.join(OUT_DIR, "finops_data.xlsx")
wb.save(xlsx_path)
print(f"Wrote {xlsx_path}")
print(f"Wrote {os.path.join(OUT_DIR, 'finops_data.json')}")
